"""
Unit tests that need no running services (no MongoDB, Ollama, Qdrant or server).

Run from the repo root:
    .venv\\Scripts\\python -m unittest backend.tests.test_units -v
"""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest import mock

import pymupdf
from docx import Document
from docx.shared import Inches
from PIL import Image

from backend.app.agent import doc_context as dc
from backend.app.auth.security import hash_password, hash_token, verify_password
from backend.app.outputs import charts, render
from backend.app.outputs.intent import detect
from backend.app.outputs.pdf_images import extract_docx_images, extract_images
from backend.app.routes import tasks as task_routes
from backend.app.tools import ocr


def _png(path: Path, width: int, height: int) -> None:
    Image.new("RGB", (width, height), "#3366cc").save(path)


class IntentTests(unittest.TestCase):
    CASES = [
        # goal, has_file, docx, pdf, image
        ("create a word file pf this summary", True, True, False, None),
        ("create a word file and a pdf of this summary", True, True, True, None),
        ("make a pdf of the summary", True, False, True, None),
        ("export the answer as PDF", True, False, True, None),
        ("give me a docx and a pdf file of key points", True, True, True, None),
        ("give me the graph image", True, False, False, "extract"),
        ("show me all the images in this pdf", True, False, False, "extract"),
        ("create a bar chart of the corrosion rates", True, False, False, "generate"),
        ("generate an image of a refinery", False, False, False, "generate"),
        # ordinary questions must NOT switch to the file pipeline
        ("summarise", True, False, False, None),
        ("summarize this pdf", True, False, False, None),
        ("create a summary of this pdf", True, False, False, None),
        ("generate a report from the uploaded pdf", True, False, False, None),
        ("What is the aim of this experiment?", True, False, False, None),
        ("explain the graph in the document", True, False, False, None),
        ("give me the word count", True, False, False, None),
    ]

    def test_cases(self):
        for goal, has_file, docx, pdf, image in self.CASES:
            with self.subTest(goal=goal):
                found = detect(goal, has_file=has_file)
                self.assertEqual((found.docx, found.pdf, found.image), (docx, pdf, image))

    def test_corrosion_prompt_is_left_to_the_agent(self):
        goal = "Read the report, calculate corrosion rate and generate an approval note."
        found = detect(goal, has_file=True)
        self.assertTrue(found.corrosion)
        self.assertFalse(found.wants_files)

    def test_explicit_formats_from_the_ui(self):
        found = detect("summarize", formats=["docx", "pdf", "image"], has_file=True)
        self.assertTrue(found.docx and found.pdf)
        # With a document attached, the box means "take its pictures" (drawing a chart is the fallback).
        self.assertEqual(found.image, "extract")
        # With nothing attached there is nothing to extract from, so a picture is drawn.
        self.assertEqual(detect("summarize", formats=["image"], has_file=False).image, "generate")


class RenderTests(unittest.TestCase):
    def test_headings_written_as_bullets_become_headings(self):
        blocks = render.parse_blocks("- # Title\n- Aim\n- Install it.\n")
        self.assertEqual(blocks[0], {"type": "title", "text": "Title"})

    def test_short_bullet_before_a_sentence_is_promoted_but_a_plain_list_is_not(self):
        text = "# Doc\n- Install MySQL Server\n- Download the installer from the website.\n- Run it.\n"
        kinds = [b["type"] for b in render.parse_blocks(text)]
        self.assertEqual(kinds, ["title", "h", "ul"])
        plain = [b["type"] for b in render.parse_blocks("# Fruit\n- Apples\n- Pears\n- Plums\n- Cherries\n")]
        self.assertEqual(plain, ["title", "ul"])

    def test_label_value_lines_stay_bullets(self):
        text = "# Doc\n- Enter the details:\n- Username: root\n- Password: the one you set during installation\n"
        self.assertNotIn("Username: root", [b.get("text") for b in render.parse_blocks(text)])

    def test_code_fences_and_tables(self):
        blocks = render.parse_blocks("# T\n```sql\nSELECT 1;\n```\n| a | b |\n|---|---|\n| 1 | 2 |\n")
        kinds = [b["type"] for b in blocks]
        self.assertIn("code", kinds)
        self.assertEqual(blocks[kinds.index("table")]["rows"], [["1", "2"]])

    def test_chatty_preface_is_removed(self):
        self.assertEqual(render.clean_model_text("Here is your document:\n\n# Real title"), "# Real title")

    def test_docx_and_pdf_are_real_files(self):
        blocks = render.parse_blocks("# Report\n## Part\n- one\n- **two**\n| x | y |\n|--|--|\n| 1 | 2 |\n")
        with tempfile.TemporaryDirectory() as tmp:
            docx_path = render.render_docx(blocks, Path(tmp) / "r.docx")
            pdf_path = render.render_pdf(blocks, Path(tmp) / "r.pdf")
            doc = Document(docx_path)
            self.assertEqual(doc.paragraphs[0].text, "Report")
            self.assertEqual(len(doc.tables), 1)
            text = "".join(page.get_text() for page in pymupdf.open(pdf_path))
            self.assertIn("Report", text)
            self.assertNotIn("**", text)


class ChartDataTests(unittest.TestCase):
    def test_common_line_formats(self):
        cases = {
            "colon+unit": ("P-101: 0.14 mm/yr\nP-103: 0.30 mm/yr", ["P-101", "P-103"], [0.14, 0.3]),
            "comma": ("P-103, 0.30\nP-105, 1,200", ["P-103", "P-105"], [0.3, 1200.0]),
            "spaced dash": ("**P-104** - 0.08\nQ3 sales - 45", ["P-104", "Q3 sales"], [0.08, 45.0]),
            "bullets": ("- Apples: 5\n* Pears: 7", ["Apples", "Pears"], [5.0, 7.0]),
        }
        for name, (text, labels, values) in cases.items():
            with self.subTest(name):
                parsed = charts.parse_chart_data(text)
                self.assertIsNotNone(parsed)
                self.assertEqual((parsed[2], parsed[3]), (labels, values))

    def test_rejects_placeholders_and_too_little_data(self):
        self.assertIsNone(charts.parse_chart_data("label: 5\nlabel: 7\nlabel: 9"))
        self.assertIsNone(charts.parse_chart_data("NONE"))
        self.assertIsNone(charts.parse_chart_data("A: 5"))

    def test_duplicate_labels_are_made_unique(self):
        parsed = charts.parse_chart_data("Q1: 5\nQ1: 7\nQ2: 9")
        self.assertEqual(len(set(parsed[2])), 3)


class LongDocumentTests(unittest.TestCase):
    def test_chunks_cover_the_text_with_overlap(self):
        text = "x" * 5000
        chunks = dc._chunks(text)
        self.assertGreater(len(chunks), 3)
        self.assertTrue(all(len(c) <= dc.CHUNK_CHARS for c in chunks))

    def test_evenly_spaced_includes_first_and_last(self):
        picked = dc._evenly_spaced(100, 6)
        self.assertEqual((picked[0], picked[-1]), (0, 99))
        self.assertEqual(dc._evenly_spaced(3, 6), [0, 1, 2])

    def test_keyword_shortlist_finds_the_matching_chunk(self):
        chunks = [f"filler text number {i} about routine housekeeping" for i in range(200)]
        chunks[137] = "the vault access phrase is ORCHID-7731-DELTA"
        picked = dc._shortlist(chunks, "What is the vault access phrase?")
        self.assertIn(137, picked)
        self.assertLessEqual(len(picked), dc.SHORTLIST_SIZE + dc.SPREAD_SAMPLE + 1)

    def test_short_text_is_returned_unchanged(self):
        self.assertEqual(dc.select_context("short", "anything"), "short")

    def test_whole_document_notes_only_for_long_summaries(self):
        long_text = "word " * 10_000
        self.assertTrue(dc.needs_whole_document_notes(long_text, "summarize this document"))
        self.assertFalse(dc.needs_whole_document_notes(long_text, "what is the launch date?"))
        self.assertFalse(dc.needs_whole_document_notes("short", "summarize this document"))

    def test_notes_cover_spread_sections_and_survive_a_failing_call(self):
        text = "\n".join(f"paragraph {i} " + "y" * 900 for i in range(120))
        calls = []

        def ask(prompt: str, max_tokens: int) -> str:
            calls.append(prompt)
            if len(calls) == 2:
                raise RuntimeError("model hiccup")
            return "- a point"

        notes, covered, total = dc.long_document_notes(text, ask)
        self.assertEqual(len(calls), dc.MAX_SECTIONS)
        self.assertEqual(covered, dc.MAX_SECTIONS - 1)
        self.assertGreater(total, dc.MAX_SECTIONS)
        self.assertIn("[Part 1 of", notes)


class FileReaderTests(unittest.TestCase):
    def test_docx_text_includes_tables_in_reading_order(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "a.docx"
            doc = Document()
            doc.add_paragraph("Before the table.")
            table = doc.add_table(rows=2, cols=2)
            for r, row in enumerate([("Asset", "Value"), ("P-77", "8.4")]):
                for c, value in enumerate(row):
                    table.rows[r].cells[c].text = value
            doc.add_paragraph("After the table.")
            doc.save(path)
            text = ocr.extract_text(str(path))
        self.assertLess(text.index("Before"), text.index("P-77 | 8.4"))
        self.assertLess(text.index("P-77 | 8.4"), text.index("After"))

    def test_plain_text_files(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "n.txt"
            path.write_text("héllo world", encoding="utf-8")
            self.assertEqual(ocr.extract_text(str(path)), "héllo world")

    def test_huge_images_are_shrunk_but_small_ones_are_not(self):
        with tempfile.TemporaryDirectory() as tmp:
            big, small = Path(tmp) / "big.png", Path(tmp) / "small.png"
            _png(big, 6000, 3000)
            _png(small, 800, 600)
            shrunk = ocr._shrunk_copy_if_huge(str(big))
            self.assertIsNotNone(shrunk)
            self.assertEqual(Image.open(shrunk).size, (3000, 1500))
            Path(shrunk).unlink()
            self.assertIsNone(ocr._shrunk_copy_if_huge(str(small)))

    def test_scanned_pdf_ocr_is_capped(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "scan.pdf"
            doc = pymupdf.open()
            for _ in range(ocr.MAX_OCR_PAGES + 15):
                doc.new_page()
            doc.save(path)
            with mock.patch.object(ocr, "_ocr_pdf_page_as_image", return_value="text") as fake:
                result = ocr.extract_text(str(path))
        self.assertEqual(fake.call_count, ocr.MAX_OCR_PAGES)
        self.assertIn(f"only the first {ocr.MAX_OCR_PAGES} of {ocr.MAX_OCR_PAGES + 15} pages", result)


class ImageExtractionTests(unittest.TestCase):
    def test_docx_images_come_out_in_document_order_and_tiny_ones_are_skipped(self):
        with tempfile.TemporaryDirectory() as tmp:
            first, second, icon = Path(tmp) / "1.png", Path(tmp) / "2.png", Path(tmp) / "icon.png"
            _png(first, 620, 320)
            _png(second, 520, 420)
            _png(icon, 24, 24)
            doc = Document()
            doc.add_picture(str(first), width=Inches(4))
            doc.add_picture(str(icon), width=Inches(0.3))
            doc.add_picture(str(second), width=Inches(4))
            docx_path = Path(tmp) / "d.docx"
            doc.save(docx_path)
            found = extract_docx_images(str(docx_path), Path(tmp) / "out")
        self.assertEqual([(f.width, f.height) for f in found], [(620, 320), (520, 420)])

    def test_pdf_header_logo_repeated_on_every_page_is_skipped(self):
        with tempfile.TemporaryDirectory() as tmp:
            logo, figure = Path(tmp) / "logo.png", Path(tmp) / "fig.png"
            _png(logo, 400, 120)
            _png(figure, 800, 500)
            doc = pymupdf.open()
            for number in range(3):
                page = doc.new_page()
                page.insert_image(pymupdf.Rect(40, 20, 240, 80), filename=str(logo))       # running header
                if number == 1:
                    page.insert_image(pymupdf.Rect(60, 300, 460, 550), filename=str(figure))  # the real figure
            pdf = Path(tmp) / "d.pdf"
            doc.save(pdf)
            found = extract_images(str(pdf), Path(tmp) / "out")
        self.assertEqual([f.page for f in found], [2])


class UploadHelperTests(unittest.TestCase):
    def test_kind_of(self):
        self.assertEqual(task_routes._kind_of("a.docx"), "docx")
        self.assertEqual(task_routes._kind_of("a.PDF"), "pdf")
        self.assertEqual(task_routes._kind_of("a.png"), "image")
        self.assertEqual(task_routes._kind_of("a.zip"), "file")

    def test_magic_bytes(self):
        self.assertTrue(task_routes._looks_like(".pdf", b"%PDF-1.7"))
        self.assertFalse(task_routes._looks_like(".pdf", b"not a pdf"))
        self.assertTrue(task_routes._looks_like(".docx", b"PK\x03\x04"))
        self.assertFalse(task_routes._looks_like(".docx", b"plain"))
        self.assertTrue(task_routes._looks_like(".txt", b"anything"))

    def test_stored_step_responses_are_capped_but_short_ones_untouched(self):
        limit = task_routes.MAX_STORED_RESPONSE
        steps = [{"step": "a", "response": "x" * (limit + 500)}, {"step": "b", "response": "short"}, {"step": "c", "response": None}]
        capped = task_routes._cap_steps(steps)
        self.assertTrue(capped[0]["response"].startswith("x" * limit))
        self.assertIn("500 more characters not stored", capped[0]["response"])
        self.assertEqual((capped[1]["response"], capped[2]["response"]), ("short", None))
        self.assertEqual(len(steps[0]["response"]), limit + 500)  # the original is not modified

    def test_json_safe_replaces_infinity(self):
        self.assertEqual(task_routes._json_safe({"a": [float("inf"), 1.5]}), {"a": [None, 1.5]})


class PasswordTests(unittest.TestCase):
    def test_hash_and_verify(self):
        stored = hash_password("correct horse")
        self.assertTrue(verify_password("correct horse", stored))
        self.assertFalse(verify_password("wrong", stored))
        self.assertNotEqual(stored, hash_password("correct horse"))  # random salt
        self.assertFalse(verify_password("x", "not-a-valid-hash"))

    def test_session_token_hash_is_stable_and_not_the_token(self):
        self.assertEqual(hash_token("abc"), hash_token("abc"))
        self.assertNotIn("abc", hash_token("abc"))


if __name__ == "__main__":
    unittest.main()
